import asyncio
import io
import unittest

from app.services.document_pipeline import (
    Block,
    build_chunks,
    is_heading,
    page_needs_ocr,
    parse_blocks,
    parse_non_pdf,
    pdf_section_blocks,
    split_sentences,
)
from app.services.token_windows import make_rerank_pairs


class PipelineTests(unittest.TestCase):
    def test_course_heading_rules_exclude_metadata_versions_and_addresses(self):
        for title in ("第02章：模型的创建与调用", "第二章 模型调用", "2.1 模型创建", "2.1.1初始化", "## 模型调用"):
            with self.subTest(title=title):
                self.assertTrue(is_heading(title))
        for text in ("讲师：尚硅谷-宋红康", "官网：尚硅谷", "2.1.0", "127.0.0.1", "3.14 是圆周率。"):
            with self.subTest(text=text):
                self.assertFalse(is_heading(text))

    def test_pdf_chapters_persist_kind_section_and_page_in_blocks(self):
        import pymupdf as fitz

        pdf = fitz.open()
        page = pdf.new_page()
        for y, text, size in (
            (50, "第02章：模型的创建与调用", 20),
            (90, "讲师：尚硅谷-宋红康", 12),
            (115, "官网：尚硅谷", 12),
            (150, "2.1 模型创建", 16),
            (180, "本节介绍如何创建模型实例。", 12),
        ):
            page.insert_text((50, y), text, fontname="china-s", fontsize=size)
        page = pdf.new_page()
        page.insert_text((50, 50), "继续说明模型的创建方法。", fontname="china-s", fontsize=12)
        page.insert_text((50, 90), "第03章：模型配置", fontname="china-s", fontsize=20)
        page.insert_text((50, 130), "这里介绍配置选项。", fontname="china-s", fontsize=12)
        data = pdf.tobytes()
        pdf.close()

        async def no_ocr(image):
            raise AssertionError("Text PDF must not invoke OCR")

        blocks = asyncio.run(parse_blocks(data, ".pdf", no_ocr))
        by_text = {block.text: block for block in blocks}
        chapter = by_text["第02章：模型的创建与调用"]
        self.assertEqual(chapter.kind, "heading")
        self.assertEqual(chapter.section, chapter.text)
        self.assertEqual(chapter.page, 1)
        self.assertEqual(by_text["讲师：尚硅谷-宋红康"].kind, "paragraph")
        self.assertEqual(by_text["官网：尚硅谷"].kind, "paragraph")
        self.assertIn("第02章", by_text["本节介绍如何创建模型实例。"].section)
        self.assertIn("2.1", by_text["继续说明模型的创建方法。"].section)
        self.assertEqual(by_text["继续说明模型的创建方法。"].page, 2)
        self.assertEqual(by_text["这里介绍配置选项。"].section, "第03章：模型配置")
        chunks = build_chunks(blocks)
        self.assertTrue(any("2.1" in c["section_title"] for c in chunks))
        self.assertFalse(any("创建模型实例" in c["content"] and "配置选项" in c["content"] for c in chunks))

    def test_pdf_font_heading_survives_chunk_grouping(self):
        import pymupdf as fitz

        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((50, 50), "Model Factory", fontsize=20)
        page.insert_text((50, 90), "This paragraph explains constructing model instances with explicit parameters.", fontsize=11)
        page.insert_text((50, 140), "Invocation", fontsize=20)
        page.insert_text((50, 180), "This paragraph explains invoking a model and reading the output response.", fontsize=11)
        data = pdf.tobytes()
        pdf.close()

        async def no_ocr(image):
            raise AssertionError("No OCR")

        blocks = asyncio.run(parse_blocks(data, ".pdf", no_ocr))
        self.assertEqual([b.text for b in blocks if b.kind == "heading"], ["Model Factory", "Invocation"])
        chunks = build_chunks(blocks)
        self.assertEqual(len(chunks), 2)
        self.assertEqual(chunks[0]["section_title"], "Model Factory")
        self.assertEqual(chunks[1]["section_title"], "Invocation")
        self.assertNotIn("output response", chunks[0]["content"])

    def test_pdf_metadata_is_not_heading_just_because_it_is_large(self):
        blocks = pdf_section_blocks(
            [[("讲师：尚硅谷-宋红康", 24)], [("官网：尚硅谷", 24)], [("普通正文。", 12)]],
            page=1, body_size=12, headings=[],
        )
        self.assertTrue(all(block.kind == "paragraph" for block in blocks))

    def test_pdf_ocr_headings_do_not_require_font_metadata(self):
        blocks = pdf_section_blocks(
            [[("第02章：模型的创建与调用", 0), ("讲师：尚硅谷-宋红康", 0)],
             [("2.1 模型创建", 0), ("普通正文。", 0)]], page=1, body_size=0, headings=[],
        )
        self.assertEqual([block.kind for block in blocks], ["heading", "paragraph", "heading", "paragraph"])
        self.assertEqual(blocks[-1].section, "第02章：模型的创建与调用 / 2.1 模型创建")

    def test_sections_do_not_mix(self):
        chunks = build_chunks([Block("项目名称：甲\n使用Python。\n项目名称：乙\n使用Java。")])
        self.assertEqual(len(chunks), 2)
        self.assertNotIn("Java", chunks[0]["content"])
        self.assertIn("项目名称：乙", chunks[1]["section_title"])

    def test_long_record_repeats_question(self):
        chunks = build_chunks([Block("标准问题：怎么申请？\n答案：" + "请先提交材料。" * 80,
                                     kind="record", sheet="FAQ", row=2)], 100, 10)
        self.assertGreater(len(chunks), 1)
        for chunk in chunks:
            self.assertTrue(chunk["content"].startswith("标准问题：怎么申请？"))
            self.assertLessEqual(len(chunk["content"]), 100)
            self.assertEqual(chunk["row"], 2)

    def test_sentence_boundary_and_tail(self):
        text = "甲" * 30 + "。" + "乙" * 30 + "。" + "尾部证据。"
        chunks = split_sentences(text, 40, 0)
        self.assertEqual("".join(chunks), text)
        self.assertTrue(chunks[0].endswith("。"))
        self.assertIn("尾部证据", chunks[-1])

    def test_invalid_sizes(self):
        for size, overlap in [(0, 0), (100, -1), (100, 100)]:
            with self.assertRaises(ValueError):
                build_chunks([Block("abc")], size, overlap)

    def test_page_ranges_follow_child_evidence(self):
        chunks = build_chunks([Block("甲" * 70 + "。", page=1), Block("乙" * 70 + "。", page=2)], 80, 0)
        self.assertEqual(chunks[0]["page_start"], 1)
        self.assertEqual(chunks[-1]["page_start"], 2)

    def test_repeated_heading_is_new_parent(self):
        chunks = build_chunks([Block("项目经历"), Block("Python"), Block("项目经历"), Block("Java")])
        self.assertEqual(len(chunks), 2)
        self.assertNotEqual(chunks[0]["parent_id"], chunks[1]["parent_id"])

    def test_xlsx_columns_and_records(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.append(["标准问题", "条件", "答案"])
        sheet.append(["怎么申请", None, "提交材料"])
        sheet.append(["怎么撤销", "未审批", "点击撤销"])
        stream = io.BytesIO()
        book.save(stream)
        blocks = parse_non_pdf(stream.getvalue(), ".xlsx")
        chunks = build_chunks(blocks)
        self.assertEqual(len(chunks), 2)
        self.assertIn("答案：提交材料", chunks[0]["content"])
        self.assertNotIn("撤销", chunks[0]["content"])
        self.assertEqual(chunks[0]["row"], 2)

    def test_docx_order(self):
        from docx import Document
        doc = Document()
        doc.add_heading("岗位要求", 1)
        doc.add_paragraph("Python")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "技能"
        table.cell(0, 1).text = "年限"
        table.cell(1, 0).text = "Java"
        table.cell(1, 1).text = "三年"
        doc.add_paragraph("末尾")
        stream = io.BytesIO()
        doc.save(stream)
        blocks = parse_non_pdf(stream.getvalue(), ".docx")
        self.assertEqual(blocks[0].text, "岗位要求")
        self.assertIn("年限", blocks[-2].text)
        self.assertEqual(blocks[-1].text, "末尾")

    def test_scan_footer_triggers_ocr(self):
        self.assertTrue(page_needs_ocr("第1页", True))
        self.assertFalse(page_needs_ocr("", False))

    def test_mixed_pdf_ocr_only_image_page(self):
        import fitz
        pdf = fitz.open()
        pdf.new_page().insert_text((50, 50), "Readable page " * 10)
        page = pdf.new_page()
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 10, 10), False)
        pix.clear_with(255)
        page.insert_image(page.rect, stream=pix.tobytes("png"))
        data = pdf.tobytes()
        pdf.close()
        calls = []

        async def ocr(image):
            calls.append(image)
            return "扫描简历中的工作经历"

        blocks = asyncio.run(parse_blocks(data, ".pdf", ocr))
        self.assertEqual(len(calls), 1)
        self.assertEqual(blocks[-1].page, 2)
        self.assertIn("工作经历", blocks[-1].text)

    def test_ocr_failure_is_not_partial_success(self):
        import fitz
        pdf = fitz.open()
        page = pdf.new_page()
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 10, 10), False)
        pix.clear_with(255)
        page.insert_image(page.rect, stream=pix.tobytes("png"))
        data = pdf.tobytes()
        pdf.close()

        async def ocr(image):
            raise RuntimeError("OCR failed")

        with self.assertRaises(RuntimeError):
            asyncio.run(parse_blocks(data, ".pdf", ocr))

    def test_token_windows_include_tail_and_fit(self):
        class Tokenizer:
            def num_special_tokens_to_add(self, pair=True): return 3

            def encode(self, text, second=None, add_special_tokens=False, **kwargs):
                return list(map(ord, text + (second or ""))) + ([0] * 3 if add_special_tokens else [])

            def decode(self, ids, **kwargs): return "".join(map(chr, ids))

        tokenizer = Tokenizer()
        pairs, owners = make_rerank_pairs(tokenizer, "问题" * 100,
                                          ["甲" * 400 + "尾部证据", "短文"], 64)
        self.assertTrue(any("尾部证据" in pair[1] for pair in pairs))
        self.assertEqual(owners[-1], 1)
        for q, text in pairs:
            self.assertLessEqual(len(tokenizer.encode(q, text, add_special_tokens=True)), 64)


if __name__ == "__main__":
    unittest.main()
