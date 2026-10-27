"""Structure-preserving parsing and chunking; no database or model imports."""

from __future__ import annotations

import asyncio
import base64
import hashlib
import io
import math
import re
import zipfile
from collections import Counter
from collections.abc import Awaitable, Callable
from dataclasses import asdict, dataclass, replace

PIPELINE_VERSION = "structure-v5-pdf-headings"


@dataclass
class Block:
    text: str
    section: str = ""
    page: int | None = None
    sheet: str | None = None
    row: int | None = None
    kind: str = "paragraph"


def heading_level(text: str) -> int | None:
    """Conservative textual heading rules, including Chinese course chapters."""
    text = text.strip()
    if not text or len(text) > 100 or "\n" in text:
        return None
    markdown = re.match(r"^(#{1,6})\s+\S", text)
    if markdown:
        return len(markdown[1])
    if not re.search(r"[。！？!?；;]", text):
        chapter = re.match(r"^第[零〇一二三四五六七八九十百千两\d]+([章节篇部])\s*[：:、.．-]?\s*\S", text)
        if chapter:
            return 2 if chapter[1] == "节" else 1
        numbered = re.match(r"^(\d{1,2}(?:[.．]\d{1,2}){1,4})(?:[ \t：:、]+|(?=[\u4e00-\u9fff]))([^\d\W].*)$", text)
        if numbered:
            return min(6, len(re.split(r"[.．]", numbered[1])))
    if len(text) < 60 and re.match(
        r"^(教育背景|教育经历|工作经历|实习经历|项目经历|项目经验|专业技能|技能清单|"
        r"个人技能|自我评价|岗位职责|工作职责|任职要求|岗位要求|加分项|福利待遇|"
        r"项目名称|公司名称|项目[：:]|Education\b|Experience\b|Skills\b|Projects\b|Requirements\b)",
        text, re.I,
    ):
        return 1
    return None


def is_heading(text: str) -> bool:
    return heading_level(text) is not None


def page_needs_ocr(text: str, has_images: bool) -> bool:
    compact = re.sub(r"\s+", "", text)
    if not compact:
        return has_images
    good = len(re.findall(r"[\w\u4e00-\u9fff]", compact))
    return (has_images and good < 40) or text.count("\ufffd") / len(compact) > 0.02 or good / len(compact) < 0.4


def order_pdf_blocks(raw: list, page_width: float) -> list:
    """Read text blocks using recursive whitespace cuts (left-to-right documents).

    A vertical gutter takes precedence over row order; full-width headings first
    need a horizontal cut. Ambiguous/overlapping layouts retain geometric order.
    This is a conservative heuristic, not a learned layout/table detector.
    """
    blocks = [block for block in raw if len(block) >= 7 and block[6] == 0 and block[4].strip()]
    if len(blocks) > 10000:
        raise ValueError("PDF 页面文本块数量超限")
    if any(not all(math.isfinite(float(value)) for value in block[:4]) for block in blocks):
        raise ValueError("PDF 页面坐标无效")
    if any(block[2] < block[0] or block[3] < block[1] for block in blocks):
        raise ValueError("PDF 页面文本块边界无效")

    def gaps(items, axis):
        intervals = sorted((item[axis], item[axis + 2]) for item in items)
        end = intervals[0][1]
        found = []
        for start, stop in intervals[1:]:
            if start > end:
                found.append((start - end, (start + end) / 2))
            end = max(end, stop)
        return sorted(found, reverse=True)

    def split(items, depth=0):
        fallback = sorted(items, key=lambda item: (item[1], item[0]))
        if len(items) < 2 or depth >= 12:
            return fallback
        # Avoid mistaking small spacing between adjacent words for a column gutter.
        for gap, cut in gaps(items, 0):
            if gap < max(12.0, page_width * 0.025):
                continue
            left = [item for item in items if item[2] < cut]
            right = [item for item in items if item[0] > cut]
            if min(len(left), len(right)) < 2:
                continue
            if min(max(item[2] for item in side) - min(item[0] for item in side) for side in (left, right)) < max(
                72.0, page_width * 0.15
            ):
                continue
            return split(left, depth + 1) + split(right, depth + 1)
        for gap, cut in gaps(items, 1):
            if gap < 3.0:
                continue
            above = [item for item in items if item[3] < cut]
            below = [item for item in items if item[1] > cut]
            if above and below:
                return split(above, depth + 1) + split(below, depth + 1)
        return fallback

    return split(blocks)


def pdf_text_layout(page) -> tuple[list[list[tuple[str, float]]], float]:
    """Extract ordered text lines and font sizes without copying image payloads."""
    import pymupdf as fitz

    layout = page.get_text("dict", sort=False, flags=fitz.TEXTFLAGS_DICT & ~fitz.TEXT_PRESERVE_IMAGES)
    raw, line_groups = [], {}
    font_weights: Counter[float] = Counter()
    for index, block in enumerate(layout.get("blocks", [])):
        if block.get("type") != 0:
            continue
        lines = []
        for line in block.get("lines", []):
            spans = line.get("spans", [])
            text = "".join(span.get("text", "") for span in spans).strip()
            if not text:
                continue
            weights: Counter[float] = Counter()
            for span in spans:
                size = round(float(span.get("size", 0)), 1)
                if not math.isfinite(size) or size <= 0:
                    continue
                weight = len(span.get("text", "").strip())
                weights[size] += weight
                font_weights[size] += weight
            size = max(sorted(weights), key=weights.get) if weights else 0.0
            lines.append((text, size))
        if not lines:
            continue
        raw.append((*block["bbox"], "\n".join(text for text, _ in lines), index, 0))
        line_groups[index] = lines
    ordered = order_pdf_blocks(raw, page.rect.width)
    body_size = max(sorted(font_weights), key=font_weights.get) if font_weights else 0.0
    return [line_groups[block[5]] for block in ordered], body_size


def pdf_section_blocks(groups: list[list[tuple[str, float]]], *, page: int,
                       body_size: float, headings: list[tuple[int, str]]) -> list[Block]:
    """Persist PDF paragraph/heading kinds and carry chapter paths across pages.

    Numbered/known titles use text rules. Other short, enlarged standalone lines
    may be headings. This does not claim to recognize PDF tables or code blocks.
    """
    out: list[Block] = []

    def section() -> str:
        return " / ".join(title for _, title in headings)

    def flush(paragraph: list[str]):
        if paragraph:
            out.append(Block("\n".join(paragraph), section=section(), page=page))
            paragraph.clear()

    for group in groups:
        paragraph: list[str] = []

        for text, size in group:
            level = heading_level(text)
            metadata = re.match(r"^(讲师|授课讲师|作者|官网|网址|邮箱|电话|课程来源)[：:]", text)
            if (level is None and not metadata and 0 < len(text) <= 60
                    and not re.search(r"[。！？!?；;]|https?://|www\.", text)
                    and body_size > 0 and size >= max(body_size * 1.25, body_size + 2)):
                level = 1 if size >= body_size * 1.6 else 2
            if level is None:
                paragraph.append(text)
                continue
            flush(paragraph)
            while headings and headings[-1][0] >= level:
                headings.pop()
            headings.append((level, text))
            out.append(Block(text, section=section(), page=page, kind="heading"))
        flush(paragraph)
    return out


class ParseGuard:
    def __init__(self, max_chars, check_deadline):
        self.max_chars = max_chars
        self.check = check_deadline
        self.used = 0

    def add(self, text):
        self.check()
        self.used += len(text)
        if self.used > self.max_chars:
            raise ValueError("文档解析文本超限，请拆分文件")


def check_archive(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 4096 or sum(item.file_size for item in entries) > 32 * 1024 * 1024:
            raise ValueError("文档解压大小超限，请拆分文件")


async def parse_blocks(
    data: bytes, extension: str, ocr: Callable[[str], Awaitable[str]], *, max_chars=200000, check_deadline=lambda: None
) -> list[Block]:
    guard = ParseGuard(max_chars, check_deadline)
    guard.check()
    if extension in (".docx", ".xlsx"):
        await asyncio.to_thread(check_archive, data)
    if extension == ".pdf":
        import pymupdf as fitz

        blocks = []
        headings: list[tuple[int, str]] = []
        with fitz.open(stream=data, filetype="pdf") as pdf:
            if len(pdf) > 500:
                raise ValueError("PDF 页数超限，请拆分文件")
            for number, page in enumerate(pdf, 1):
                guard.check()
                groups, body_size = await asyncio.to_thread(pdf_text_layout, page)
                texts = ["\n".join(line for line, _ in group) for group in groups]
                text = "\n\n".join(texts)
                if page_needs_ocr(text, bool(page.get_images())):
                    if page.rect.width * page.rect.height * (150 / 72) ** 2 > 20_000_000:
                        raise ValueError("PDF 页面渲染大小超限")
                    png = await asyncio.to_thread(lambda page=page: page.get_pixmap(dpi=150).tobytes("png"))
                    text = await ocr(base64.b64encode(png).decode("ascii"))
                    if not text.strip() or page_needs_ocr(text, False):
                        # 纯图片页（如海报/大图）OCR 也提不出可靠文字：
                        # 跳过该页并继续，不让单页拖垮整个文档。
                        blocks.append(Block(f"[第 {number} 页为图片页，未提取到可靠文本]", page=number))
                        continue
                    groups = [[(line.strip(), 0.0) for line in part.splitlines() if line.strip()]
                              for part in text.split("\n\n") if part.strip()]
                    body_size = 0.0
                guard.add(text)
                blocks.extend(pdf_section_blocks(groups, page=number, body_size=body_size, headings=headings))
        return blocks
    return await asyncio.to_thread(parse_non_pdf, data, extension, guard)


def parse_markdown(text: str, guard: ParseGuard) -> list[Block]:
    """Keep Markdown source and heading paths without rendering HTML or fetching URLs.

    Recognize ATX/Setext headings outside fenced code. Paragraphs retain lists,
    tables and inline markup verbatim; code fences are kept as separate records.
    """
    lines = text.replace("\r\n", "\n").replace("\r", "\n").splitlines()
    blocks: list[Block] = []
    headings: list[tuple[int, str]] = []
    buffer: list[str] = []
    fence: tuple[str, int] | None = None

    def section() -> str:
        return " / ".join(title for _, title in headings)

    def flush(kind: str = "markdown") -> None:
        if buffer:
            content = "\n".join(buffer).strip()
            if content:
                if len(blocks) >= 10000:
                    raise ValueError("Markdown 段落数量超限，请拆分文件")
                blocks.append(Block(content, section=section(), kind=kind))
            buffer.clear()

    index = 0
    while index < len(lines):
        guard.check()
        line = lines[index]
        if fence:
            buffer.append(line)
            closing = re.fullmatch(r" {0,3}(`{3,}|~{3,})[ \t]*", line)
            if closing and closing[1][0] == fence[0] and len(closing[1]) >= fence[1]:
                flush("record")
                fence = None
            index += 1
            continue

        opening = re.match(r"^ {0,3}(`{3,}|~{3,})(.*)$", line)
        if opening and (opening[1][0] != "`" or "`" not in opening[2]):
            flush()
            fence = (opening[1][0], len(opening[1]))
            buffer.append(line)
            index += 1
            continue

        heading = re.match(r"^ {0,3}(#{1,6})(?:[ \t]+|$)(.*)$", line)
        level = 0
        title = ""
        consumed = 1
        if heading:
            level = len(heading[1])
            title = re.sub(r"[ \t]+#+[ \t]*$", "", heading[2]).strip()
        elif line.strip() and index + 1 < len(lines):
            underline = re.fullmatch(r" {0,3}(=+|-+)[ \t]*", lines[index + 1])
            if underline and not re.match(r"^\s*(?:>|[-+*]\s|\d+[.)]\s|\|)", line):
                level = 1 if underline[1][0] == "=" else 2
                title = line.strip()
                consumed = 2
        if level and title:
            flush()
            while headings and headings[-1][0] >= level:
                headings.pop()
            headings.append((level, title))
            buffer.append("\n".join(lines[index : index + consumed]))
            flush()
        elif not line.strip():
            flush()
        else:
            buffer.append(line)
        index += consumed

    # An unclosed fence remains code through end-of-file.
    flush("record" if fence else "markdown")
    return blocks


def parse_non_pdf(data: bytes, extension: str, guard=None) -> list[Block]:
    guard = guard or ParseGuard(200000, lambda: None)
    if extension == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        try:
            if len(book.sheetnames) > 50:
                raise ValueError("工作表数量超限")
            for sheet in book:
                guard.check()
                if (sheet.max_row or 0) > 10000 or (sheet.max_column or 0) > 200:
                    raise ValueError("工作表行列数量超限，请拆分文件")
                headers = None
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    guard.check()
                    if number > 10000 or len(row) > 200:
                        raise ValueError("工作表行列数量超限")
                    if not any(v is not None for v in row):
                        continue
                    if headers is None:
                        headers = [str(v).strip() if v is not None else f"列{i + 1}" for i, v in enumerate(row)]
                        guard.add(" ".join(headers))
                        continue
                    fields = [
                        f"{headers[i] if i < len(headers) else f'列{i + 1}'}：{v}"
                        for i, v in enumerate(row)
                        if v is not None and str(v).strip()
                    ]
                    text = "\n".join(fields)
                    guard.add(text)
                    out.append(Block(text, section=sheet.title, sheet=sheet.title, row=number, kind="record"))
        finally:
            book.close()
        return out
    if extension == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = Document(io.BytesIO(data))
        out = []
        section = ""
        for count, child in enumerate(doc.element.body.iterchildren(), 1):
            guard.check()
            if count > 10000:
                raise ValueError("文档段落数量超限")
            if child.tag.endswith("}p"):
                paragraph = Paragraph(child, doc)
                text = paragraph.text.strip()
                if not text:
                    continue
                guard.add(text)
                if (paragraph.style and paragraph.style.name.startswith("Heading")) or is_heading(text):
                    section = text
                out.append(Block(text, section=section))
            elif child.tag.endswith("}tbl"):
                table = Table(child, doc)
                header = ""
                for index, row in enumerate(table.rows):
                    guard.check()
                    if index >= 10000 or len(row.cells) > 200:
                        raise ValueError("文档表格行列数量超限")
                    text = " | ".join(cell.text.strip() for cell in row.cells)
                    if index:
                        text = header + "\n" + text
                    else:
                        header = text
                    guard.add(text)
                    out.append(Block(text, section=section, kind="record"))
        return out
    if extension in (".txt", ".md"):
        encodings = ("utf-16",) if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "gb18030")
        for encoding in encodings:
            try:
                text = data.decode(encoding)
            except UnicodeError:
                continue
            guard.add(text)
            if extension == ".md":
                return parse_markdown(text, guard)
            return [Block(p.strip()) for p in re.split(r"\n\s*\n", text.replace("\r\n", "\n")) if p.strip()]
        raise ValueError("无法识别文本编码")
    raise ValueError(f"不支持的文件格式：{extension}")


def split_sentences(text: str, size: int, overlap: int) -> list[str]:
    """Prefer paragraph/newline/sentence boundaries; character split is the last resort."""
    if size < 1 or not 0 <= overlap < size:
        raise ValueError("要求 size > overlap >= 0")
    units = []
    for unit in re.split(r"(?<=[。！？!?；;])|(?<=\n)", text):
        if unit:
            units.extend(unit[i : i + size] for i in range(0, len(unit), size))
    out, current = [], ""
    for unit in units:
        if current and len(current) + len(unit) > size:
            out.append(current.strip())
            # Overlap only whole trailing sentences, and only when they fit.
            trailing = re.split(r"(?<=[。！？!?；;])|(?<=\n)", current)
            tail = ""
            for part in reversed(trailing):
                if len(part + tail) > overlap:
                    break
                tail = part + tail
            current = tail if len(tail + unit) <= size else ""
        current += unit
    if current.strip():
        out.append(current.strip())
    return out


def build_chunks(blocks: list[Block], size: int = 800, overlap: int = 100) -> list[dict]:
    if size < 64 or not 0 <= overlap < size:
        raise ValueError("要求 size >= 64 且 0 <= overlap < size")
    normalized = []
    for block in blocks:
        if block.kind == "heading":
            normalized.append(block)
            continue
        if block.kind in ("record", "markdown"):
            normalized.append(block)
            continue
        buffer = []
        active_section = block.section
        for line in block.text.splitlines():
            if is_heading(line.strip()):
                if buffer:
                    normalized.append(replace(block, text="\n".join(buffer), section=active_section))
                    buffer = []
                active_section = line.strip("# ")
                normalized.append(replace(block, text=line.strip(), section=active_section))
            else:
                buffer.append(line)
        if buffer and "\n".join(buffer).strip():
            normalized.append(replace(block, text="\n".join(buffer), section=active_section))
    blocks = normalized
    parents: list[tuple[str, list[Block]]] = []
    section = ""
    for block in blocks:
        if block.kind == "heading" or is_heading(block.text):
            section = block.text.strip("# ")
        section = block.section or section
        # Spreadsheet/table rows must not merge into neighbouring records.
        if (
            block.kind == "record"
            or block.kind == "heading" or is_heading(block.text)
            or not parents
            or parents[-1][0] != section
            or parents[-1][1][-1].kind == "record"
        ):
            parents.append((section, [block]))
        else:
            parents[-1][1].append(block)
    chunks = []
    for parent_index, (title, group) in enumerate(parents):
        parent_text = "\n\n".join(b.text for b in group)
        spans, offset = [], 0
        for block in group:
            spans.append((offset, offset + len(block.text), block.page))
            offset += len(block.text) + 2
        prefix = f"章节：{title}\n" if title else ""
        if group[0].kind == "record":
            # Repeat a FAQ question for long answers, without inventing content.
            question = next(
                (line for line in parent_text.splitlines() if re.match(r"^(标准问题|问题|question)[：:]", line, re.I)),
                "",
            )
            if question:
                prefix += question + "\n"
        prefix = prefix[: min(160, size // 3)]
        pieces = split_sentences(parent_text, size - len(prefix), min(overlap, max(0, size - len(prefix) - 1)))
        previous_start = -1
        for piece in pieces:
            content = prefix + piece
            start = parent_text.find(piece, previous_start + 1)
            if start < 0:
                start = parent_text.find(piece)
            previous_start = start
            pages = [
                page for left, right, page in spans if page is not None and left < start + len(piece) and right > start
            ]
            chunks.append(
                {
                    "content": content,
                    "chunk_index": len(chunks),
                    "content_hash": hashlib.sha256(content.encode()).hexdigest(),
                    "parent_id": str(parent_index),
                    "section_title": title,
                    "page_start": min(pages) if pages else None,
                    "page_end": max(pages) if pages else None,
                    "sheet": group[0].sheet,
                    "row": group[0].row,
                    "parent_text": parent_text,
                }
            )
    return chunks


def serialize_blocks(blocks: list[Block]) -> list[dict]:
    return [asdict(block) for block in blocks]
